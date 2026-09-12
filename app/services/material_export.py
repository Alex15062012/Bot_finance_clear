"""Экспорт материалов админки в PDF и Word."""

from __future__ import annotations

import io
import re
from pathlib import Path

from app.db.models import Material

_FONTS_DIR = Path(__file__).resolve().parent.parent / "admin" / "static" / "fonts"
_DEJAVU = _FONTS_DIR / "DejaVuSans.ttf"
_DEJAVU_BOLD = _FONTS_DIR / "DejaVuSans-Bold.ttf"

_FONT_CANDIDATES = [
    _DEJAVU,
    Path(r"C:\Windows\Fonts\arial.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
]


def material_plain_parts(material: Material) -> list[tuple[str, str]]:
    """Секции для экспорта: (заголовок, текст)."""
    parts: list[tuple[str, str]] = []
    if material.title:
        parts.append(("title", material.title.strip()))
    if material.description:
        parts.append(("description", material.description.strip()))
    if material.content:
        parts.append(("content", material.content.strip()))
    if material.url:
        parts.append(("url", material.url.strip()))
    return parts


def safe_filename(title: str, ext: str) -> str:
    base = re.sub(r"[^\w\s\-а-яА-ЯёЁ]+", "", title or "material", flags=re.UNICODE)
    base = re.sub(r"\s+", "_", base.strip())[:80] or "material"
    return f"{base}.{ext.lstrip('.')}"


def content_disposition(filename: str, material_id: int, ext: str) -> str:
    from urllib.parse import quote

    ascii_name = f"material_{material_id}.{ext.lstrip('.')}"
    return (
        f'attachment; filename="{ascii_name}"; '
        f"filename*=UTF-8''{quote(filename)}"
    )


def _resolve_font() -> Path:
    for path in _FONT_CANDIDATES:
        if path.is_file():
            return path
    raise FileNotFoundError(
        "Не найден шрифт с кириллицей для PDF. "
        "Положите DejaVuSans.ttf в app/admin/static/fonts/ "
        "или установите системный Arial/DejaVu."
    )


def export_material_docx(material: Material) -> bytes:
    from docx import Document

    doc = Document()
    for kind, text in material_plain_parts(material):
        if kind == "title":
            doc.add_heading(text, level=1)
        elif kind == "url":
            doc.add_paragraph(f"Ссылка: {text}")
        else:
            doc.add_paragraph(text)

    if not material_plain_parts(material):
        doc.add_paragraph("(пустой материал)")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_material_pdf(material: Material) -> bytes:
    from fpdf import FPDF

    font_path = _resolve_font()
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.add_font("Body", fname=str(font_path))
    bold_path = _DEJAVU_BOLD if _DEJAVU_BOLD.is_file() else font_path
    pdf.add_font("Body", style="B", fname=str(bold_path))

    parts = material_plain_parts(material)
    if not parts:
        pdf.set_font("Body", size=12)
        pdf.multi_cell(0, 8, "(пустой материал)")
    for kind, text in parts:
        if kind == "title":
            pdf.set_font("Body", style="B", size=16)
            pdf.multi_cell(0, 10, text)
            pdf.ln(4)
        elif kind == "url":
            pdf.set_font("Body", size=11)
            pdf.multi_cell(0, 7, f"Ссылка: {text}")
            pdf.ln(2)
        else:
            pdf.set_font("Body", size=12)
            pdf.multi_cell(0, 7, text)
            pdf.ln(3)

    out = pdf.output()
    if isinstance(out, (bytes, bytearray)):
        return bytes(out)
    return str(out).encode("latin-1")
